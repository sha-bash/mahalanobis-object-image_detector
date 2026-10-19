"""Two-page manuscript; numerical tables are read from frozen v4 evidence."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.platypus import TableStyle

import build_learned_report
from build_learned_report import block, docx, pdf, ST

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'experiments/evidence/instance_v4'
OUT = ROOT / 'output/article_instance_20261011'


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def number(value, digits=2):
    return f'{value:.{digits}f}'.replace('.', ',')


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rows = {r['method']: r for r in read(BASE / 'test_summary.json') if r['delta'] == .02}
    video = {r['method']: r for r in read(ROOT / 'experiments/evidence/davis_v4/summary.json') if r['subset'] == 'main_eval'}
    names = ['visual_baseline', 'visual_strict', 'gate_text_s2027_full']
    labels = ['DINOv2 исходный', 'DINOv2 строгий', 'DINOv2 + текстовый фильтр']
    stats = read(BASE / 'statistics.json')
    assert (rows[names[1]]['tp'], rows[names[1]]['fp'], rows[names[2]]['tp'], rows[names[2]]['fp']) == (134, 119, 128, 110)
    B = []

    def add(kind, text=None, **kw):
        B.append(block(kind, text, **kw))

    add('title', 'Текстовая проверка ложных совпадений при поиске экземпляра объекта по одному изображению')
    add('center', 'В. О. Шабашев')
    add('center', 'Ульяновский государственный технический университет, Ульяновск, Россия')
    add('p', '<b>Аннотация.</b> Исследуется текстовая проверка кандидатов при поиске конкретного экземпляра по одному эталону. Сопоставлены визуальный метод, ужесточение его порога и фильтр на основе описаний, обучаемого проектора и расстояния Махаланобиса. На новых экземплярах стоимость ошибок снизилась на 6,28%, однако потеря полноты превысила заданный допуск. Дополнительно проверен перенос на видео без настройки порогов.')
    add('small', 'Ключевые слова: поиск экземпляра, визуально-языковая модель, метрическое обучение, расстояние Махаланобиса, ложные принятия.')
    add('head', '1 Постановка задачи')
    add('p', 'Даны одно эталонное изображение q конкретного экземпляра и кадр I. Требуется выдать рамку этого же экземпляра либо отказ, в том числе при наличии похожих объектов. Совпадения категории недостаточно. При обработке видео решение принимается для каждого кадра без заданной начальной рамки. Обучение общей модели на других экземплярах допускается; дополнительные изображения тестовой цели не используются.')
    add('p', 'Цель исследования - уменьшить ложные принятия при ограниченной потере полноты. Для N решений определим риск J=(5FP+FN)/N, полноту R=TP/(TP+FN) и FPR=FP/(FP+TN), где TP - верное принятие, FP - ложное принятие, FN - пропуск, TN - верный отказ. Рабочая гипотеза: добавление текста снижает J и FPR относительно ужесточения визуального порога при потере R не более 0,02 относительно исходного метода. Нулевая гипотеза отрицает наличие такого преимущества. Отказ на всех входах исключается ограничением полноты.')
    add('head', '2 Метод текстовой проверки')
    add('p', 'Визуальная оценка - косинусное расстояние признаков DINOv2-small [1]. Qwen3.5:0.8b независимо описывает эталон и кандидата; all-MiniLM-L6-v2 семейства Sentence-BERT [2] кодирует тексты. Авторский MLP [3] размерности 384→256→128→64 с BatchNorm, GELU и dropout адаптирован к сопоставлению новых экземпляров. Обучение: batch-hard triplet, L2-нормировка выхода, margin 0,2, Adam 0,001, 30 эпох. Для нормированных выходов z используется')
    add('image', path=ROOT / 'experiments/evidence/instance_v3/mahalanobis_formula.png', height=36)
    add('p', 'где a=z(q), b=z(c), а общая ковариация Ledoit-Wolf [4] оценивается по симметризованным разностям обучающих пар одного экземпляра. Текст проверяет лишь визуально принятые неоднозначные пары; пороги выбираются на валидации. Сравнение с косинусом на тех же MLP-признаках отделяет влияние метрики. Вклад работы - адаптация проектора и проверка дополнительной пользы текста против простого ужесточения визуального решения.')
    add('head', '3 Экспериментальный протокол')
    add('p', 'Из Washington RGB-D [5] до вычисления оценок выбраны 80 новых физических экземпляров 20 категорий: по 40 для валидации и теста, всего 400 изображений. Единственный эталон и четыре кандидата взяты из разных видео одного стенда. Тест содержит 320 пар, поровну совпадающих и различных экземпляров одной категории. Обучение на 184 других экземплярах (920 изображений) не пересекается с новым разбиением.')
    add('p', 'Исходный порог максимизирует средний F1 на валидации: R₀=91,25%. Строгий визуальный метод и текстовый фильтр минимизируют J при R≥R₀−0,02=89,25%. Пороги и основной seed 2027 фиксируются до теста; 2028 и 2029 проверяют устойчивость. Подтверждение требует снижения FPR, положительной нижней границы 95% интервала выигрыша J, нижней границы изменения R не ниже −0,02 и тестовой R≥89,25%. Парный bootstrap проводится по 20 категориям, 10 000 повторов; пары не считаются независимыми.')

    add('break')
    add('head', '4 Результаты и обсуждение')
    add('small', 'Таблица 1. Сопоставление вырезок на новых экземплярах. R и FPR усреднены по категориям; TP и FP - суммарные счётчики.')
    table = [['Метод', 'TP', 'FP', 'R, %', 'FPR, %', '5FP+FN']]
    for name, label in zip(names, labels):
        r = rows[name]
        table.append([label, r['tp'], r['fp'], number(100*r['tpr']), number(100*r['fpr']), 5*r['fp']+r['fn']])
    add('table', rows=table, widths=[200,40,40,65,65,65])
    add('p', 'Относительно строгого визуального контроля текстовый фильтр уменьшил FP со 119 до 110, но TP снизилось со 134 до 128 (табл. 1). Отклонены 13 ложных совпадений контроля и приняты четыре других: чистое уменьшение FP равно девяти. Стоимость ошибок снизилась с 621 до 582, или на 6,28%. Полнота уменьшилась на 3,75 п.п., что превышает допуск 2 п.п.')
    add('p', f'Снижение риска на пару равно {number(stats["risk_reduction_vs_strict"],4)}; 95% интервал [{number(stats["risk_ci95"][0],4)}; {number(stats["risk_ci95"][1],4)}] включает ноль. Интервал изменения полноты [-7,50; -0,625] п.п. не подтверждает её сохранение с заданным допуском. Даже исходная тестовая R=83,75% ниже валидационного ориентира 89,25%, что указывает на ограниченную переносимость калибровки.')
    add('p', 'Махаланобис относительно косинуса на том же MLP дал на пять FP меньше, но потерял ещё три TP. Стоимость ошибок для seeds 2027/2028/2029 составила 582/637/582; один запуск оказался хуже строгого контроля (621). Следовательно, устойчивое превосходство метрики не установлено. Среди 138 пар, направленных на текстовую проверку, отклонены 13 из 80 чужих экземпляров и 6 из 58 правильных совпадений.')
    add('p', 'Отдельно использованы 20 заранее выбранных видео DAVIS 2017 [6]: 29 целей, 232 решения. Эталон вырезан из первого кадра, исключённого из оценки; это протокол той же сцены. YOLOv8n формирует общие кандидаты, DINOv2 выбирает ближайший. Пороги переносятся из RGB-D без настройки. TP требует IoU≥0,5; неверная рамка при наличии цели даёт FP и FN, пропуск имеет IoU=0.')
    add('small', 'Таблица 2. Локализация на видео. R и IoU на кадрах наличия цели усреднены по последовательностям; счётчики - по 232 решениям.')
    table = [['Метод', 'TP', 'FP', 'R, %', 'Средний IoU']]
    for name, label in zip(names, labels):
        r = video[name]
        table.append([label, r['tp'], r['fp'], number(100*r['macro']['recall']), number(r['macro']['mean_iou'],3)])
    add('table', rows=table, widths=[225,45,45,75,85])
    add('p', 'На видео относительно строгого контроля получено на три FP меньше при потере одного TP (табл. 2). Однако эффект сосредоточен в двух последовательностях; случаев отсутствия видимой цели всего четыре. Зависимость сцен и возможное пересечение с предобучением неизвестны. Запуски выполнены последовательно на CPU с 8 ГБ RAM; работа в реальном времени не заявляется.')
    add('head', '5 Выводы')
    add('p', 'Текстовая проверка способна отсекать часть ложных совпадений сверх ужесточения визуального порога. В проведённом эксперименте наблюдается снижение стоимости ошибок, однако преимущество при потере полноты не более 2 п.п. не подтверждено. Результат характеризует измеренный компромисс, а не универсальное улучшение поиска. Дальнейшая проверка требует новых независимых съёмок, эталонов из другой сцены и большего числа эпизодов отсутствия и возвращения цели.')
    add('head', 'Литература')
    refs = [
        '1. Oquab M. et al. DINOv2: Learning Robust Visual Features without Supervision. 2023. arXiv:2304.07193. https://arxiv.org/abs/2304.07193.',
        '2. Reimers N., Gurevych I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks // EMNLP-IJCNLP. 2019. P. 3982-3992. DOI: 10.18653/v1/D19-1410.',
        '3. Шабашев В. О. Mahalanobis Concept Drift Detector. Программная реализация MLP. https://github.com/sha-bash/mahalanobis-concept-drift. Версия fb3bccd (дата обращения: 11.10.2026).',
        '4. Ledoit O., Wolf M. A well-conditioned estimator for large-dimensional covariance matrices // Journal of Multivariate Analysis. 2004. Vol. 88, no. 2. P. 365-411. DOI: 10.1016/S0047-259X(03)00096-4.',
        '5. Lai K., Bo L., Ren X., Fox D. A Large-Scale Hierarchical Multi-View RGB-D Object Dataset // ICRA. 2011. https://rgbd-dataset.cs.washington.edu/dataset.html.',
        '6. Pont-Tuset J. et al. The 2017 DAVIS Challenge on Video Object Segmentation. 2017. arXiv:1704.00675. https://arxiv.org/abs/1704.00675.',
    ]
    for ref in refs:
        add('small', ref)
    ST['PaperBody'].fontSize = 10.4
    ST['PaperBody'].leading = 12.1
    ST['PaperBody'].spaceAfter = 4
    ST['SmallRU'].fontSize = 8.4
    ST['SmallRU'].leading = 9.5
    ST['SmallRU'].spaceAfter = 3
    original_table = build_learned_report.table
    def bordered_table(rows, widths=None):
        t = original_table(rows, widths)
        t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.HexColor('#D9D9D9')),
                               ('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
        return t
    build_learned_report.table = bordered_table
    pdf(B, OUT/'article.pdf', paper=True)
    docx(B, OUT/'article.docx', paper=True)
    d = Document(OUT/'article.docx')
    d.core_properties.title = B[0]['text']
    d.core_properties.subject = 'Контролируемая оценка текстового фильтра при поиске экземпляра'
    normal = d.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    d.styles['Caption'].font.size = Pt(8.4)
    d.styles['Caption'].paragraph_format.space_after = Pt(3)
    d.styles['Caption'].paragraph_format.line_spacing = 1.0
    for p in d.paragraphs:
        if p.style.name == 'Title' or p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p.style.name == 'Normal' and p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t in d.tables:
        t.style = 'Normal Table'
        borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            element = OxmlElement('w:'+edge)
            for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:
                element.set(qn('w:'+k),v)
            borders.append(element)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i == 0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'),'E8EDF1')
                    cell._tc.get_or_add_tcPr().append(shading)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = i==0
        following = t._tbl.getnext()
        if following is not None and following.tag == qn('w:p'):
            from docx.text.paragraph import Paragraph
            Paragraph(following, d._body).paragraph_format.space_before = Pt(5)
    d.save(OUT/'article.docx')
    pages = len(PdfReader(OUT/'article.pdf').pages)
    print({'pdf_pages':pages, 'blocks':len(B)})
    assert pages==2, 'Adjust manuscript layout before delivery'


if __name__ == '__main__':
    main()
